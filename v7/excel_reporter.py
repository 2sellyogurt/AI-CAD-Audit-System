# -*- coding: utf-8 -*-
"""协调单生成器 —— v5 JSON → Excel协调单"""
import json, os
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelReporter:
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    A_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    B_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    C_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    THIN_BORDER = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    def __init__(self, conflicts_json_path: str):
        with open(conflicts_json_path, "r", encoding="utf-8") as f:
            self.conflicts = json.load(f)

    def generate(self, output_path: str = "coordination_sheet.xlsx"):
        wb = Workbook()
        self._write_summary(wb)
        self._write_detail(wb)
        wb.save(output_path)
        print(f"协调单已生成: {output_path}")
        return output_path

    def _write_summary(self, wb):
        ws = wb.active
        ws.title = "汇总统计"
        ws.merge_cells("A1:F1")
        ws["A1"] = "AI智能审图系统 v5 —— 冲突协调单"
        ws["A1"].font = Font(name="微软雅黑", size=16, bold=True)

        sev = Counter(c.get("severity", "") for c in self.conflicts)
        typ = Counter(c.get("type", "") for c in self.conflicts)

        ws["A3"] = "严重程度分布"
        ws["A3"].font = Font(bold=True)
        row = 4
        for s in ("A", "B", "C"):
            ws.cell(row=row, column=1, value=s)
            ws.cell(row=row, column=2, value=sev.get(s, 0))
            row += 1

        ws["A8"] = "类型分布"
        ws["A8"].font = Font(bold=True)
        row = 9
        for t, cnt in typ.most_common():
            ws.cell(row=row, column=1, value=t)
            ws.cell(row=row, column=2, value=cnt)
            row += 1

        ws["A14"] = f"总冲突簇: {len(self.conflicts)}"
        ws["A15"] = f"跨专业冲突: {sum(1 for c in self.conflicts if c.get('type') != 'egress_width')}"
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 12

    def _write_detail(self, wb):
        ws = wb.create_sheet("冲突明细")
        headers = ["#", "严重度", "类型", "楼层", "描述", "风险", "涉及图纸", "规范条文", "重复数"]
        col_widths = [5, 8, 22, 8, 50, 40, 35, 30, 8]

        for i, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.THIN_BORDER
            ws.column_dimensions[get_column_letter(i)].width = col_widths[i - 1]

        sev_order = {"A": 0, "B": 1, "C": 2}
        sorted_conflicts = sorted(self.conflicts, key=lambda c: (sev_order.get(c.get("severity", "Z"), 99), c.get("type", "")))

        for idx, c in enumerate(sorted_conflicts, 1):
            row = idx + 1
            severity = c.get("severity", "")
            fill = {"A": self.A_FILL, "B": self.B_FILL, "C": self.C_FILL}.get(severity)

            values = [
                idx, severity, c.get("type", ""), c.get("floor", ""),
                c.get("description", "")[:200], c.get("risk", "")[:150],
                c.get("involved", "")[:120], c.get("std", "")[:100],
                c.get("duplicate_count", 1)
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=(col >= 5))
                if fill and col <= 3:
                    cell.fill = fill

        ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    import sys
    json_path = sys.argv[1] if len(sys.argv) > 1 else "spatial_conflicts_v5.json"
    output = sys.argv[2] if len(sys.argv) > 2 else "coordination_sheet_v5.xlsx"
    reporter = ExcelReporter(json_path)
    reporter.generate(output)