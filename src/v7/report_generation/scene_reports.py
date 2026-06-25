# -*- coding: utf-8 -*-
"""
场景报告生成器 - 共享模块
从 spatial_conflicts + UAI审查数据生成三种场景报告：
  场景1_基础错漏排查报告
  场景2_跨专业一致性校验报告  
  场景3_强条合规性审查报告
"""

import json
import os
from collections import defaultdict
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

CONFLICT_CN = {
    "beam_duct_overlap": "梁-风管重叠",
    "beam_pipe_overlap": "梁-管线重叠",
    "column_pipe_conflict": "柱-管线冲突",
    "pipe_crossing": "管线交叉",
    "duct_through_wall": "风管穿墙",
    "egress_width": "疏散宽度不足",
}

DISCIPLINE_CN = {
    "building": "建筑", "structure": "结构", "plumbing": "给排水",
    "hvac": "暖通", "electrical": "电气", "fire": "消防",
    "facade": "幕墙", "landscape": "景观", "foundation": "基坑",
    "interior": "装饰",
}

SEVERITY_CN = {
    "A": "A(强条)", "B": "B(一般)", "C": "C(标注不全)", "D": "D(建议)",
}


def setup_doc(doc):
    """设置文档基础样式"""
    if not HAS_DOCX:
        return
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def add_section_title(doc, title, level=1):
    """添加章节标题"""
    if not HAS_DOCX:
        return
    heading = doc.add_heading(title, level=level)
    heading.style.font.name = '微软雅黑'
    heading.style.font.size = Pt(12 if level == 1 else 11)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT


def generate_scene1_report(output_dir, issues, spatial_conflicts, project_name="项目"):
    """生成场景1：基础错漏排查报告"""
    if not HAS_DOCX:
        return None
    
    doc = Document()
    setup_doc(doc)
    
    t = doc.add_heading(f"基础错漏排查报告 v7.0 — LLM代理审查", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"项目名称: {project_name}")
    doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()
    
    add_section_title(doc, "一、审查概况")
    doc.add_paragraph(f"审查总数: {len(issues)} 条")
    
    severity_counts = defaultdict(int)
    for issue in issues:
        severity_counts[issue.get("severity", "unknown")] += 1
    
    for sev, cnt in severity_counts.items():
        doc.add_paragraph(f"  {SEVERITY_CN.get(sev, sev)}: {cnt} 条")
    
    return doc


def generate_scene2_report(output_dir, issues, spatial_conflicts, project_name="项目"):
    """生成场景2：跨专业一致性校验报告"""
    if not HAS_DOCX:
        return None
    
    doc = Document()
    setup_doc(doc)
    
    t = doc.add_heading(f"跨专业一致性校验报告 v7.0 — LLM代理审查", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"项目名称: {project_name}")
    doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()
    
    add_section_title(doc, "一、跨专业冲突分析")
    doc.add_paragraph(f"空间冲突总数: {len(spatial_conflicts)}")
    
    return doc


def generate_scene3_report(output_dir, issues, spatial_conflicts, project_name="项目"):
    """生成场景3：强条合规审查报告"""
    if not HAS_DOCX:
        return None
    
    doc = Document()
    setup_doc(doc)
    
    t = doc.add_heading(f"强条合规审查报告 v7.0 — LLM代理审查", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"项目名称: {project_name}")
    doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()
    
    add_section_title(doc, "一、强条合规分析")
    critical_issues = [i for i in issues if i.get("severity") == "A"]
    doc.add_paragraph(f"强条问题总数: {len(critical_issues)} 条")
    
    return doc


def generate_all_scenes(output_dir, issues, spatial_conflicts, project_name="项目"):
    """生成所有三个场景报告"""
    results = {}
    
    scene1 = generate_scene1_report(output_dir, issues, spatial_conflicts, project_name)
    if scene1:
        path1 = os.path.join(output_dir, "场景1_基础错漏排查报告.docx")
        scene1.save(path1)
        results["scene1"] = path1
    
    scene2 = generate_scene2_report(output_dir, issues, spatial_conflicts, project_name)
    if scene2:
        path2 = os.path.join(output_dir, "场景2_跨专业一致性校验报告.docx")
        scene2.save(path2)
        results["scene2"] = path2
    
    scene3 = generate_scene3_report(output_dir, issues, spatial_conflicts, project_name)
    if scene3:
        path3 = os.path.join(output_dir, "场景3_强条合规性审查报告.docx")
        scene3.save(path3)
        results["scene3"] = path3
    
    return results